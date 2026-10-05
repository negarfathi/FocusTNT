from pathlib import Path
import argparse
import math

from openpyxl import load_workbook
from scipy.stats import binomtest, wilcoxon


RQ_SHEET = "RQ6-2"

ANALYZERS = [
    "Athena",
    "PROTON",
    "UAutomizer",
    "AProVE",
    "CPAchecker",
    "2LS",
]

ROW_MAP = {
    "Slice": dict(zip(ANALYZERS, range(7, 13))),
    "Cncrt": dict(zip(ANALYZERS, range(13, 19))),
    "Slice+Cncrt": dict(zip(ANALYZERS, range(19, 25))),
    "Cncrt+Slice": dict(zip(ANALYZERS, range(25, 31))),
}

BASE = {
    "correct": "P",
    "t_ratio": "R",
    "nt_ratio": "T",
    "tvt": "W",
}

CONFIGS = {
    "Slice": {
        "correct": "AI",
        "t_ratio": "AK",
        "nt_ratio": "AM",
        "tvt": "AP",
    },
    "Cncrt": {
        "correct": "BB",
        "t_ratio": "BD",
        "nt_ratio": "BF",
        "tvt": "BI",
    },
    "Slice+Cncrt": {
        "correct": "BU",
        "t_ratio": "BW",
        "nt_ratio": "BY",
        "tvt": "CB",
    },
    "Cncrt+Slice": {
        "correct": "CN",
        "t_ratio": "CP",
        "nt_ratio": "CR",
        "tvt": "CU",
    },
}


def to_float(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)

    s = str(value).strip()
    if not s:
        return None

    try:
        if s.endswith("%"):
            return float(s[:-1]) / 100.0
        return float(s)
    except ValueError:
        return None


def is_correct(value):
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() == "correct"


def solved(correct_value, ratio_value, task):
    if not is_correct(correct_value):
        return False

    ratio = to_float(ratio_value)
    if ratio is None:
        return False

    if task == "NT":
        return ratio > 0.0

    if task == "T":
        return math.isclose(ratio, 1.0, rel_tol=0.0, abs_tol=1e-12)

    raise ValueError(task)


def exact_mcnemar_p(recoveries, losses):
    n = recoveries + losses
    if n == 0:
        return 1.0
    return float(
        binomtest(recoveries, n=n, p=0.5, alternative="two-sided").pvalue
    )


def holm_adjust(pvalues):
    m = len(pvalues)
    order = sorted(range(m), key=lambda i: pvalues[i])

    adjusted = [None] * m
    running_max = 0.0

    for rank, idx in enumerate(order):
        multiplier = m - rank
        candidate = min(1.0, multiplier * float(pvalues[idx]))
        running_max = max(running_max, candidate)
        adjusted[idx] = min(1.0, running_max)

    return adjusted


def safe_wilcoxon(base_values, cfg_values):
    n = len(base_values)

    if n == 0:
        return 0, 0.0, 1.0

    differences = [c - b for b, c in zip(base_values, cfg_values)]

    if all(
        math.isclose(d, 0.0, rel_tol=0.0, abs_tol=1e-12)
        for d in differences
    ):
        return n, 0.0, 1.0

    result = wilcoxon(
        base_values,
        cfg_values,
        zero_method="wilcox",
        alternative="two-sided",
        method="auto",
    )

    return n, float(result.statistic), float(result.pvalue)


def original_program_rows(ws):
    rows = []

    for r in range(3, ws.max_row + 1):
        name = ws[f"A{r}"].value
        gt = ws[f"F{r}"].value

        if name not in (None, "") and gt in ("NT", "T"):
            rows.append(r)

    return rows


def choose_output_path(input_path: Path, requested_output=None):
    if requested_output:
        output_path = Path(requested_output).expanduser().resolve()
    else:
        output_path = input_path.with_name(
            input_path.stem + "_rq6_2_filled.xlsx"
        )

    # Never overwrite the input file.
    if output_path == input_path:
        output_path = input_path.with_name(
            input_path.stem + "_rq6_2_filled.xlsx"
        )

    # Never overwrite an existing output file either.
    if output_path.exists():
        counter = 2
        base = output_path.with_suffix("")
        suffix = output_path.suffix

        while True:
            candidate = Path(f"{base}_{counter}{suffix}")
            if not candidate.exists():
                output_path = candidate
                break
            counter += 1

    return output_path


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Compute RQ6-2 from an Excel workbook and save the results "
            "to a NEW Excel file."
        )
    )

    parser.add_argument(
        "input_excel",
        help="Path to the source .xlsx file",
    )

    parser.add_argument(
        "-o",
        "--output",
        help=(
            "Optional output .xlsx path. If omitted, a new file named "
            "<input>_rq6_2_filled.xlsx is created beside the source file."
        ),
    )

    args = parser.parse_args()

    input_path = Path(args.input_excel).expanduser().resolve()

    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    if input_path.suffix.lower() != ".xlsx":
        raise ValueError("Input file must be an .xlsx workbook.")

    output_path = choose_output_path(input_path, args.output)

    # data_only=True reads the calculated values from source-sheet formulas.
    wb_values = load_workbook(
        input_path,
        data_only=True,
        read_only=False,
    )

    # data_only=False preserves formulas/styles in the workbook copy.
    wb_write = load_workbook(
        input_path,
        data_only=False,
        read_only=False,
    )

    if RQ_SHEET not in wb_write.sheetnames:
        raise KeyError(f"Missing sheet: {RQ_SHEET}")

    for analyzer in ANALYZERS:
        if analyzer not in wb_values.sheetnames:
            raise KeyError(f"Missing analyzer sheet: {analyzer}")

    results = {}

    for analyzer in ANALYZERS:
        ws = wb_values[analyzer]
        rows = original_program_rows(ws)
        results[analyzer] = {}

        for config_name, cfg in CONFIGS.items():
            cfg_result = {}

            for task in ("NT", "T"):
                gt_rows = [
                    r for r in rows
                    if ws[f"F{r}"].value == task
                ]

                base_ratio_col = (
                    BASE["nt_ratio"]
                    if task == "NT"
                    else BASE["t_ratio"]
                )

                cfg_ratio_col = (
                    cfg["nt_ratio"]
                    if task == "NT"
                    else cfg["t_ratio"]
                )

                recoveries = 0
                losses = 0

                base_tvt = []
                cfg_tvt = []

                for r in gt_rows:
                    base_solved = solved(
                        ws[f'{BASE["correct"]}{r}'].value,
                        ws[f"{base_ratio_col}{r}"].value,
                        task,
                    )

                    cfg_solved = solved(
                        ws[f'{cfg["correct"]}{r}'].value,
                        ws[f"{cfg_ratio_col}{r}"].value,
                        task,
                    )

                    if (not base_solved) and cfg_solved:
                        recoveries += 1
                    elif base_solved and (not cfg_solved):
                        losses += 1

                    b = to_float(
                        ws[f'{BASE["tvt"]}{r}'].value
                    )
                    c = to_float(
                        ws[f'{cfg["tvt"]}{r}'].value
                    )

                    if (
                        base_solved
                        and cfg_solved
                        and b is not None
                        and c is not None
                    ):
                        base_tvt.append(b)
                        cfg_tvt.append(c)

                mcnemar_raw = exact_mcnemar_p(
                    recoveries,
                    losses,
                )

                n_pairs, w_stat, wilcoxon_raw = safe_wilcoxon(
                    base_tvt,
                    cfg_tvt,
                )

                cfg_result[task] = {
                    "recoveries": recoveries,
                    "losses": losses,
                    "mcnemar_raw": mcnemar_raw,
                    "N": n_pairs,
                    "W": w_stat,
                    "wilcoxon_raw": wilcoxon_raw,
                }

            results[analyzer][config_name] = cfg_result

    config_order = [
        "Slice",
        "Cncrt",
        "Slice+Cncrt",
        "Cncrt+Slice",
    ]

    # Holm correction over all four comparisons,
    # separately for each analyzer, task, and test family.
    for analyzer in ANALYZERS:
        for task in ("NT", "T"):
            mc_raw = [
                results[analyzer][cfg][task]["mcnemar_raw"]
                for cfg in config_order
            ]

            wc_raw = [
                results[analyzer][cfg][task]["wilcoxon_raw"]
                for cfg in config_order
            ]

            mc_adj = holm_adjust(mc_raw)
            wc_adj = holm_adjust(wc_raw)

            for i, cfg in enumerate(config_order):
                results[analyzer][cfg][task][
                    "mcnemar_holm"
                ] = mc_adj[i]

                results[analyzer][cfg][task][
                    "wilcoxon_holm"
                ] = wc_adj[i]

    rq = wb_write[RQ_SHEET]

    # Fill only the numeric result area D:O.
    for cfg in config_order:
        for analyzer in ANALYZERS:
            row = ROW_MAP[cfg][analyzer]

            nt = results[analyzer][cfg]["NT"]
            t = results[analyzer][cfg]["T"]

            values = {
                "D": nt["recoveries"],
                "E": nt["losses"],
                "F": nt["mcnemar_holm"],
                "G": t["recoveries"],
                "H": t["losses"],
                "I": t["mcnemar_holm"],
                "J": nt["N"],
                "K": nt["W"],
                "L": nt["wilcoxon_holm"],
                "M": t["N"],
                "N": t["W"],
                "O": t["wilcoxon_holm"],
            }

            for col, value in values.items():
                rq[f"{col}{row}"] = value

    for row in range(7, 31):
        for col in ("D", "E", "G", "H", "J", "M"):
            rq[f"{col}{row}"].number_format = "0"

        for col in ("K", "N"):
            rq[f"{col}{row}"].number_format = "0.0###"

        for col in ("F", "I", "L", "O"):
            rq[f"{col}{row}"].number_format = "0.000000"

    # Save ONLY to the new path.
    wb_write.save(output_path)

    print(f"Source workbook kept unchanged:")
    print(f"  {input_path}")
    print()
    print(f"New workbook created:")
    print(f"  {output_path}")


if __name__ == "__main__":
    main()
